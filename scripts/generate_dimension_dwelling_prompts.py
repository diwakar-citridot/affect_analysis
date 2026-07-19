#!/usr/bin/env python3
"""
Generate dimension-dwelling lived-experience prompts for every concept in MySQL.

Reads (dimension, concept) slug pairs from ``svarupa_concepts`` (joined to
``svarupa_dimensions``) together with, for each concept:

  * the **aspect** (``svarupa_concepts.description``),
  * the **coordinate** (``svarupa_concepts.coordinate`` JSON — seat/guṇa/valence/...),
  * the authored **per-perspective descriptions per status** from
    ``svarupa_concept_descriptions`` (``deficiency`` / ``balance`` / ``excess``).

It then builds **one prompt per (dimension, concept, status)**: each prompt combines
all perspective descriptions for that single status, plus the coordinate and aspect,
and asks for 20 lived-experience statements (10 Western + 10 Asian) in that state.
The concept-specific Step 2 ✗/✓ pairs and ``## 6. Worked Examples`` are generated
once per concept and reused across its statuses. Prompts are written as
``<dimension>__<concept>__<status>.txt`` under ``prompts/dd-mm-yyyy/``.

Concepts that have only an aspect (no authored descriptions) fall back to generating
all three states from the aspect and coordinate alone.

Usage
-----
  .venv/bin/python scripts/generate_dimension_dwelling_prompts.py
  .venv/bin/python scripts/generate_dimension_dwelling_prompts.py --dimension pancha_koshas
  .venv/bin/python scripts/generate_dimension_dwelling_prompts.py --concept annamaya_kosha
  .venv/bin/python scripts/generate_dimension_dwelling_prompts.py --output-dir ./prompts/26-05-2026
  .venv/bin/python scripts/generate_dimension_dwelling_prompts.py --log-file ./logs/prompt-generation.log
  .venv/bin/python scripts/generate_dimension_dwelling_prompts.py --no-log-file

  # Also generate the lived-experience statements for each concept via the LLM and
  # store them in an Excel workbook (requires per-concept prompt files):
  .venv/bin/python scripts/generate_dimension_dwelling_prompts.py --generate-lived-experience
  .venv/bin/python scripts/generate_dimension_dwelling_prompts.py --concept annamaya_kosha \
      --generate-lived-experience --lived-experience-output ./prompts/26-05-2026/lived_experience.xlsx

Environment (loaded from ./.env if present)
-------------------------------------------
  SVARUPA_MYSQL_HOST
  SVARUPA_MYSQL_PORT              (default: 3306)
  SVARUPA_MYSQL_USER
  SVARUPA_MYSQL_PASSWORD
  SVARUPA_MYSQL_DATABASE_MASTER   (default: svarupa_assistant_v1)
  SVARUPA_MYSQL_DATABASE          (fallback when MASTER is unset)
  BEDROCK_MODEL_ID / SVARUPA_DIMENSION_DWELLING_BEDROCK_MODEL_ID
                                  (default: us.anthropic.claude-opus-4-7)
  AWS_REGION / SVARUPA_AWS_REGION (default: us-west-2)
"""
from __future__ import annotations

import argparse
import json
import logging
import os
import sys
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date
from pathlib import Path

import pymysql

_SCRIPTS_DIR = Path(__file__).resolve().parent
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

from dimension_dwelling_prompt_template import (
    LIVED_EXPERIENCE_MAX_TOKENS,
    STATUS_CODE_TO_LABEL,
    STATUS_ORDER,
    ConceptRecord,
    build_status_prompt,
    build_step2_examples,
    build_target_quality_examples,
    display_name,
    first_sentence,
    generate_lived_experience_text,
    get_step2_bedrock_client,
    infer_profile,
    parse_lived_experience_json,
    reset_step2_bedrock_client,
    split_statements,
)

DEFAULT_MAX_WORKERS = 6

REPO_ROOT = Path(__file__).resolve().parents[1]
_SRC_DIR = REPO_ROOT / "src"
if str(_SRC_DIR) not in sys.path:
    sys.path.insert(0, str(_SRC_DIR))

from svarupa_affect.infrastructure.config import load_dotenv as _load_repo_dotenv

_load_repo_dotenv()

logger = logging.getLogger(__name__)
_LOG_FORMAT = "%(asctime)s  %(levelname)-8s  %(message)s"


def _setup_logging(*, verbose: bool, log_file: Path | None) -> None:
    level = logging.DEBUG if verbose else logging.INFO
    formatter = logging.Formatter(_LOG_FORMAT)

    root = logging.getLogger()
    root.handlers.clear()
    root.setLevel(level)

    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setLevel(level)
    console_handler.setFormatter(formatter)
    root.addHandler(console_handler)

    if log_file is not None:
        log_file.parent.mkdir(parents=True, exist_ok=True)
        file_handler = logging.FileHandler(log_file, encoding="utf-8")
        file_handler.setLevel(level)
        file_handler.setFormatter(formatter)
        root.addHandler(file_handler)
        logger.info("Logging to file: %s", log_file)

    for name in (logger.name, "dimension_dwelling_prompt_template"):
        logging.getLogger(name).setLevel(level)

    logger.setLevel(level)


def _log_concept_progress(
    *,
    index: int,
    total: int,
    record: ConceptRecord,
    status_label: str,
    prompt_chars: int,
    output_path: Path | None,
    combined_only: bool,
) -> None:
    message = (
        f"[{index}/{total}] Wrote prompt | dimension={record.dimension!r} "
        f"concept={record.concept!r} status={status_label!r} chars={prompt_chars}"
    )
    if output_path is not None:
        message += f" output={output_path}"
    elif combined_only:
        message += " output=(combined file only)"
    logger.info(message)


def _env(name: str, default: str = "") -> str:
    return (os.getenv(name) or default).strip()


def _cfg() -> dict[str, str | int]:
    database = (
        _env("SVARUPA_MYSQL_DATABASE_MASTER")
        or _env("SVARUPA_MYSQL_DATABASE")
        or "svarupa_assistant_v1"
    )
    return {
        "host": _env("SVARUPA_MYSQL_HOST", "127.0.0.1"),
        "port": int(_env("SVARUPA_MYSQL_PORT", "3306")),
        "user": _env("SVARUPA_MYSQL_USER", "root"),
        "password": _env("SVARUPA_MYSQL_PASSWORD"),
        "database": database,
    }


def _table_exists(cur: pymysql.cursors.DictCursor, db: str, table: str) -> bool:
    cur.execute(
        """
        SELECT 1
        FROM information_schema.TABLES
        WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s
        LIMIT 1
        """,
        (db, table),
    )
    return cur.fetchone() is not None


def _column_exists(cur: pymysql.cursors.DictCursor, db: str, table: str, column: str) -> bool:
    cur.execute(
        """
        SELECT 1
        FROM information_schema.COLUMNS
        WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s AND COLUMN_NAME = %s
        LIMIT 1
        """,
        (db, table, column),
    )
    return cur.fetchone() is not None


def _concept_overview_sql(cur: pymysql.cursors.DictCursor, db: str) -> str:
    """Best-effort concept-level overview from ``svarupa_concepts`` columns."""
    parts: list[str] = []
    for col in ("description",):
        if _column_exists(cur, db, "svarupa_concepts", col):
            parts.append(f"NULLIF(TRIM(c.`{col}`), '')")
    if not parts:
        return "NULL"
    if len(parts) == 1:
        return parts[0]
    return f"COALESCE({', '.join(parts)})"


def _description_pick_sql(db: str) -> str:
    """Pick one triplet description per concept (perspective, then status sort order)."""
    return f"""
    SELECT
      cd.concept_id,
      NULLIF(TRIM(cd.description), '') AS picked_description,
      ROW_NUMBER() OVER (
        PARTITION BY cd.concept_id
        ORDER BY
          CASE cd.perspective
            WHEN 'non_spiritual' THEN 0
            WHEN 'physical_body' THEN 1
            WHEN 'vital' THEN 2
            WHEN 'mental_body' THEN 3
            WHEN 'psychic' THEN 4
            WHEN 'spoken_languages' THEN 5
            WHEN 'overview' THEN 6
            ELSE 7
          END,
          st.sort_order,
          cd.description_id
      ) AS rn
    FROM `{db}`.`svarupa_concept_descriptions` cd
    JOIN `{db}`.`svarupa_status` st ON st.status_id = cd.status_id
    WHERE NULLIF(TRIM(cd.description), '') IS NOT NULL
    """


def _fetch_concepts(
    cur: pymysql.cursors.DictCursor,
    cfg: dict[str, str | int],
    *,
    dimension: str | None,
    concept: str | None,
    allowed_dimension_slugs: frozenset[str] | None = None,
) -> list[ConceptRecord]:
    """Load one row per logical concept with the best available source meaning."""
    db = str(cfg["database"])
    if not _table_exists(cur, db, "svarupa_concepts") or not _table_exists(
        cur, db, "svarupa_dimensions"
    ):
        raise RuntimeError(
            f"Expected svarupa_concepts and svarupa_dimensions in database {db!r}. "
            "Apply sql/001_svarupa_dimensions_concepts.sql or run the companion migration."
        )

    overview_expr = _concept_overview_sql(cur, db)
    has_descriptions = _table_exists(cur, db, "svarupa_concept_descriptions") and _table_exists(
        cur, db, "svarupa_status"
    )

    if has_descriptions:
        meaning_expr = f"COALESCE({overview_expr}, picked.picked_description)"
        description_join = f"""
    LEFT JOIN (
      {_description_pick_sql(db)}
    ) picked ON picked.concept_id = c.concept_id AND picked.rn = 1"""
    else:
        meaning_expr = overview_expr
        description_join = ""

    filters: list[str] = [f"{meaning_expr} IS NOT NULL"]
    params: list[str] = []
    if dimension:
        filters.append("d.slug = %s")
        params.append(dimension)
    if allowed_dimension_slugs is not None:
        if not allowed_dimension_slugs:
            return []
        placeholders = ", ".join("%s" for _ in allowed_dimension_slugs)
        filters.append(f"d.slug IN ({placeholders})")
        params.extend(sorted(allowed_dimension_slugs))
    if concept:
        filters.append("c.slug = %s")
        params.append(concept)

    has_coordinate = _column_exists(cur, db, "svarupa_concepts", "coordinate")
    coordinate_expr = "c.coordinate" if has_coordinate else "NULL"

    where_sql = " AND ".join(filters)
    sql = f"""
    SELECT
      c.concept_id AS concept_id,
      d.slug AS dimension,
      NULLIF(TRIM(d.name), '') AS db_dimension_name,
      NULLIF(TRIM(d.sanskrit_term), '') AS db_dimension_sanskrit_term,
      c.slug AS concept,
      NULLIF(TRIM(c.name), '') AS db_display_name,
      NULLIF(TRIM(c.sanskrit_term), '') AS db_sanskrit_term,
      NULLIF(TRIM(c.category), '') AS db_category,
      {overview_expr} AS aspect,
      {coordinate_expr} AS coordinate,
      {meaning_expr} AS source_meaning
    FROM `{db}`.`svarupa_concepts` c
    JOIN `{db}`.`svarupa_dimensions` d ON d.dimension_id = c.dimension_id{description_join}
    WHERE {where_sql}
    ORDER BY d.slug, c.slug
    """
    cur.execute(sql, params)
    rows = cur.fetchall()

    concept_ids = [int(row["concept_id"]) for row in rows]
    facets_by_concept = (
        _fetch_status_facets(cur, db, concept_ids) if (has_descriptions and concept_ids) else {}
    )

    records: list[ConceptRecord] = []
    for row in rows:
        meaning = (row.get("source_meaning") or "").strip()
        aspect = (row.get("aspect") or "").strip()
        if not meaning and not aspect:
            continue
        records.append(
            ConceptRecord(
                dimension=str(row["dimension"]),
                concept=str(row["concept"]),
                source_meaning=meaning or aspect,
                db_dimension_name=(row.get("db_dimension_name") or "").strip(),
                db_dimension_sanskrit_term=(row.get("db_dimension_sanskrit_term") or "").strip(),
                db_display_name=(row.get("db_display_name") or "").strip(),
                db_sanskrit_term=(row.get("db_sanskrit_term") or "").strip(),
                db_category=(row.get("db_category") or "").strip(),
                aspect=aspect,
                coordinate=_parse_coordinate(row.get("coordinate")),
                status_facets=facets_by_concept.get(int(row["concept_id"]), {}),
            )
        )
    return records


def _parse_coordinate(raw: object) -> dict[str, str]:
    """Parse the ``svarupa_concepts.coordinate`` JSON column into a flat dict of strings."""
    if raw is None:
        return {}
    if isinstance(raw, dict):
        data = raw
    else:
        try:
            data = json.loads(raw)
        except (TypeError, ValueError):
            return {}
    if not isinstance(data, dict):
        return {}
    return {str(k): str(v).strip() for k, v in data.items() if v is not None and str(v).strip()}


def _fetch_status_facets(
    cur: pymysql.cursors.DictCursor,
    db: str,
    concept_ids: list[int],
) -> dict[int, dict[str, tuple[tuple[str, str], ...]]]:
    """Load authored per-perspective descriptions grouped by concept then status code."""
    if not concept_ids:
        return {}
    placeholders = ", ".join("%s" for _ in concept_ids)
    cur.execute(
        f"""
        SELECT
          cd.concept_id AS concept_id,
          st.code AS status_code,
          NULLIF(TRIM(cd.perspective), '') AS perspective,
          NULLIF(TRIM(cd.description), '') AS description
        FROM `{db}`.`svarupa_concept_descriptions` cd
        JOIN `{db}`.`svarupa_status` st ON st.status_id = cd.status_id
        WHERE cd.concept_id IN ({placeholders})
          AND NULLIF(TRIM(cd.description), '') IS NOT NULL
        ORDER BY cd.concept_id, st.sort_order, cd.perspective
        """,
        concept_ids,
    )
    grouped: dict[int, dict[str, list[tuple[str, str]]]] = {}
    for row in cur.fetchall():
        cid = int(row["concept_id"])
        status_code = str(row["status_code"])
        perspective = (row.get("perspective") or "").strip() or "overview"
        description = (row.get("description") or "").strip()
        if not description:
            continue
        grouped.setdefault(cid, {}).setdefault(status_code, []).append((perspective, description))
    return {
        cid: {status: tuple(facets) for status, facets in by_status.items()}
        for cid, by_status in grouped.items()
    }


def _default_output_dir(base: Path | None = None) -> Path:
    root = base or Path.cwd()
    stamp = date.today().strftime("%d-%m-%Y")
    return root / "prompts" / stamp


def _safe_filename(dimension: str, concept: str, status_label: str | None = None) -> str:
    raw = f"{dimension}__{concept}"
    if status_label:
        raw += f"__{status_label}"
    return "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in raw)


def _statuses_for_record(record: ConceptRecord) -> list[str]:
    """Canonical status codes to generate for a concept.

    Prefer the statuses that actually have authored descriptions (in canonical
    order); fall back to all three states when a concept has only an aspect.
    """
    if record.status_facets:
        ordered = [code for code in STATUS_ORDER if code in record.status_facets]
        extra = [code for code in record.status_facets if code not in STATUS_ORDER]
        return ordered + sorted(extra)
    return list(STATUS_ORDER)


def _build_concept_prompts(
    record: ConceptRecord,
    *,
    siblings: list[tuple[str, str]],
    model_id: str | None,
    region: str | None,
) -> list[tuple[str, str]]:
    """Return ``[(status_code, prompt_text), ...]`` for a concept.

    Makes the one concept-scoped Step 2 LLM call, then assembles a prompt per status.
    Runs inside the thread pool, so it does no shared-state mutation or file I/O.
    """
    profile = infer_profile(record)
    step2_examples = build_step2_examples(profile, record, model_id=model_id, region=region)
    target_quality_examples = build_target_quality_examples(profile)
    prompts: list[tuple[str, str]] = []
    for status_code in _statuses_for_record(record):
        prompts.append(
            (
                status_code,
                build_status_prompt(
                    record,
                    status_code,
                    step2_examples=step2_examples,
                    target_quality_examples=target_quality_examples,
                    siblings=siblings,
                ),
            )
        )
    return prompts


def _write_prompts(
    records: list[ConceptRecord],
    output_dir: Path,
    *,
    combined_only: bool,
    model_id: str | None,
    region: str | None,
    max_workers: int = DEFAULT_MAX_WORKERS,
) -> tuple[Path, list[tuple[ConceptRecord, str, Path]], int, int]:
    """Build one prompt per (concept, status).

    The ✗/✓ calibration (Step 2) and worked examples are concept-scoped, so they
    are generated once per concept and reused across that concept's statuses. The
    Step 2 LLM calls run concurrently; file writing stays serial and ordered.
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    written: list[tuple[ConceptRecord, str, Path]] = []
    combined_parts: list[str] = []
    total = sum(len(_statuses_for_record(record)) for record in records)
    failed = 0
    succeeded = 0
    index = 0

    # Sibling contrast set per dimension: (concept_slug, display_name, gloss).
    siblings_by_dim: dict[str, list[tuple[str, str, str]]] = {}
    for rec in records:
        name = rec.db_display_name or display_name(rec.concept)
        gloss = first_sentence(rec.aspect) if rec.aspect else ""
        siblings_by_dim.setdefault(rec.dimension, []).append((rec.concept, name, gloss))

    def _siblings_for(record: ConceptRecord) -> list[tuple[str, str]]:
        return [
            (name, gloss)
            for (slug, name, gloss) in siblings_by_dim.get(record.dimension, [])
            if slug != record.concept
        ]

    # Concept-level Step 2 calls run in parallel; results keyed by record index so
    # files are still written in deterministic (dimension, concept) order.
    get_step2_bedrock_client(region=region)  # pre-init client to avoid lazy-init races
    results: dict[int, list[tuple[str, str]] | Exception] = {}
    with ThreadPoolExecutor(max_workers=max_workers) as pool:
        future_to_idx = {
            pool.submit(
                _build_concept_prompts,
                record,
                siblings=_siblings_for(record),
                model_id=model_id,
                region=region,
            ): i
            for i, record in enumerate(records)
        }
        for future in as_completed(future_to_idx):
            i = future_to_idx[future]
            try:
                results[i] = future.result()
            except Exception as exc:  # noqa: BLE001 - recorded per concept
                results[i] = exc

    for i, record in enumerate(records):
        statuses = _statuses_for_record(record)
        result = results.get(i)
        if isinstance(result, Exception) or result is None:
            failed += len(statuses)
            index += len(statuses)
            logger.error(
                "Failed to build concept-level examples; skipping all statuses | "
                "dimension=%r concept=%r statuses=%s error=%s",
                record.dimension,
                record.concept,
                statuses,
                result,
            )
            continue

        for status_code, prompt in result:
            index += 1
            status_label = STATUS_CODE_TO_LABEL.get(status_code, status_code)
            succeeded += 1
            header = (
                f"Dimension: {record.dimension}\n"
                f"Concept: {record.concept}\n"
                f"Status: {status_label}\n"
                f"{'=' * 72}\n\n"
            )
            full_text = header + prompt
            combined_parts.append(full_text)

            output_path: Path | None = None
            if not combined_only:
                filename = (
                    _safe_filename(record.dimension, record.concept, status_label) + ".txt"
                )
                output_path = output_dir / filename
                output_path.write_text(full_text, encoding="utf-8")
                written.append((record, status_code, output_path))

            _log_concept_progress(
                index=index,
                total=total,
                record=record,
                status_label=status_label,
                prompt_chars=len(prompt),
                output_path=output_path,
                combined_only=combined_only,
            )

    combined_path = output_dir / "all_prompts.txt"
    divider = f"\n\n{'#' * 72}\n\n"
    combined_text = divider.join(combined_parts)
    combined_path.write_text(combined_text, encoding="utf-8")
    logger.info(
        "Wrote combined prompts file: %s (%s prompt(s), %s chars, %s failed)",
        combined_path,
        len(combined_parts),
        len(combined_text),
        failed,
    )
    return combined_path, written, succeeded, failed


_PROMPT_HEADER_DIVIDER = f"{'=' * 72}\n\n"


def _read_prompt_text(path: Path) -> str:
    """Strip the ``Dimension/Concept`` header written by ``_write_prompts`` and return the prompt body."""
    content = path.read_text(encoding="utf-8")
    idx = content.find(_PROMPT_HEADER_DIVIDER)
    if idx == -1:
        return content
    return content[idx + len(_PROMPT_HEADER_DIVIDER) :]


def _lived_experience_ckpt_key(dimension: str, concept: str, status_label: str) -> str:
    return f"{dimension}\t{concept}\t{status_label}"


def _load_lived_experience_checkpoint(path: Path) -> dict[str, dict]:
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _save_lived_experience_checkpoint(path: Path, data: dict) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


LIVED_EXPERIENCE_COLUMNS = [
    "dimension",
    "concept",
    "display_name",
    "regional_perspective",
    "status",
    "statement_number",
    "statement",
]
# Backward-compatible alias.
_LIVED_EXPERIENCE_COLUMNS = LIVED_EXPERIENCE_COLUMNS


def _require_openpyxl():
    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for --generate-lived-experience. "
            "Install with: pip install openpyxl"
        ) from exc
    return openpyxl


def _write_lived_experience_excel(
    items: list[tuple[ConceptRecord, str, Path]],
    checkpoint: dict[str, dict],
    output_path: Path,
) -> int:
    rows: list[dict[str, str | int]] = []
    for record, status_code, _ in items:
        status_label = STATUS_CODE_TO_LABEL.get(status_code, status_code)
        entry = checkpoint.get(
            _lived_experience_ckpt_key(record.dimension, record.concept, status_label), {}
        )
        if entry.get("status") != "completed":
            continue
        display_name_val = entry.get("display_name") or infer_profile(record).display_name

        if "statements" in entry:
            # New JSON format: list of {statement, status, regional_perspective}
            for stmt_number, item in enumerate(entry["statements"], start=1):
                rows.append(
                    {
                        "dimension": record.dimension,
                        "concept": record.concept,
                        "display_name": display_name_val,
                        "regional_perspective": item["regional_perspective"],
                        "status": item["status"],
                        "statement_number": stmt_number,
                        "statement": item["statement"],
                    }
                )
        else:
            # Legacy format: flat western/asian text blocks without status tags
            for perspective, block in (
                ("western", entry.get("western_perspective", "")),
                ("asian", entry.get("asian_perspective", "")),
            ):
                for statement_number, statement in enumerate(split_statements(block), start=1):
                    rows.append(
                        {
                            "dimension": record.dimension,
                            "concept": record.concept,
                            "display_name": display_name_val,
                            "regional_perspective": perspective,
                            "status": "",
                            "statement_number": statement_number,
                            "statement": statement,
                        }
                    )

    openpyxl = _require_openpyxl()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "lived_experience"
    ws.append(_LIVED_EXPERIENCE_COLUMNS)
    for row in rows:
        ws.append([row[col] for col in _LIVED_EXPERIENCE_COLUMNS])
    wb.save(output_path)
    return len(rows)


def _generate_lived_experiences(
    items: list[tuple[ConceptRecord, str, Path]],
    *,
    output_path: Path,
    checkpoint_path: Path,
    resume: bool,
    model_id: str | None,
    region: str | None,
    max_tokens: int,
    max_workers: int = DEFAULT_MAX_WORKERS,
) -> tuple[int, int]:
    """Read each per-(concept, status) prompt file, run it through the LLM, and write Excel.

    Generation calls run concurrently; the shared checkpoint is guarded by a lock and
    saved after each completion so the run stays resumable if interrupted.
    """
    checkpoint = _load_lived_experience_checkpoint(checkpoint_path) if resume else {}
    total = len(items)
    lock = threading.Lock()
    counters = {"succeeded": 0, "failed": 0, "done": 0}

    # Split into already-completed (resumed) and to-do.
    todo: list[tuple[ConceptRecord, str, Path]] = []
    for record, status_code, prompt_path in items:
        status_label = STATUS_CODE_TO_LABEL.get(status_code, status_code)
        key = _lived_experience_ckpt_key(record.dimension, record.concept, status_label)
        if checkpoint.get(key, {}).get("status") == "completed":
            counters["succeeded"] += 1
        else:
            todo.append((record, status_code, prompt_path))
    if counters["succeeded"]:
        logger.info("Resumed %s completed lived-experience item(s) from checkpoint", counters["succeeded"])

    get_step2_bedrock_client(region=region)  # pre-init client to avoid lazy-init races

    def _work(item: tuple[ConceptRecord, str, Path]) -> None:
        record, status_code, prompt_path = item
        status_label = STATUS_CODE_TO_LABEL.get(status_code, status_code)
        key = _lived_experience_ckpt_key(record.dimension, record.concept, status_label)
        prompt_text = _read_prompt_text(prompt_path)
        try:
            raw = generate_lived_experience_text(
                prompt_text,
                model_id=model_id,
                region=region,
                max_tokens=max_tokens,
                extra={
                    "dimension": record.dimension,
                    "concept": record.concept,
                    "status": status_label,
                    "display_name": infer_profile(record).display_name,
                    "prompt_path": str(prompt_path),
                },
            )
            statements = parse_lived_experience_json(raw)
        except Exception as exc:  # noqa: BLE001 - recorded per item
            with lock:
                counters["failed"] += 1
                counters["done"] += 1
                checkpoint[key] = {"status": "error", "error": str(exc)}
                _save_lived_experience_checkpoint(checkpoint_path, checkpoint)
                done = counters["done"]
            logger.error(
                "[%s/%s] Failed to generate lived experience | dimension=%r concept=%r status=%r error=%s",
                done,
                len(todo),
                record.dimension,
                record.concept,
                status_label,
                exc,
            )
            return

        # This prompt targets a single state; coerce every statement's status to the
        # target label so a model mislabel can't leak a wrong state into the workbook.
        for stmt in statements:
            stmt["status"] = status_label
        western_count = sum(1 for s in statements if s["regional_perspective"] == "western")
        asian_count = len(statements) - western_count
        with lock:
            counters["succeeded"] += 1
            counters["done"] += 1
            checkpoint[key] = {
                "status": "completed",
                "display_name": infer_profile(record).display_name,
                "statements": statements,
            }
            _save_lived_experience_checkpoint(checkpoint_path, checkpoint)
            done = counters["done"]
        logger.info(
            "[%s/%s] Generated lived experience | dimension=%r concept=%r status=%r "
            "total=%s western=%s asian=%s",
            done,
            len(todo),
            record.dimension,
            record.concept,
            status_label,
            len(statements),
            western_count,
            asian_count,
        )

    if todo:
        with ThreadPoolExecutor(max_workers=max_workers) as pool:
            list(pool.map(_work, todo))

    rows_written = _write_lived_experience_excel(items, checkpoint, output_path)
    logger.info(
        "Wrote lived-experience workbook: %s (%s row(s)); %s succeeded, %s failed of %s total",
        output_path,
        rows_written,
        counters["succeeded"],
        counters["failed"],
        total,
    )
    return counters["succeeded"], counters["failed"]


def run_dimension_dwelling_generation(
    *,
    output_dir: Path | None = None,
    dimension: str | None = None,
    concept: str | None = None,
    allowed_dimension_slugs: frozenset[str] | None = None,
    limit: int | None = None,
    combined_only: bool = False,
    generate_lived_experience: bool = False,
    lived_experience_output: Path | None = None,
    model_id: str | None = None,
    region: str | None = None,
    lived_experience_model_id: str | None = None,
    lived_experience_max_tokens: int = LIVED_EXPERIENCE_MAX_TOKENS,
    max_workers: int = DEFAULT_MAX_WORKERS,
    resume: bool = True,
    verbose: bool = False,
    log_file: Path | None = None,
    no_log_file: bool = False,
    setup_logging: bool = True,
) -> tuple[int, Path | None]:
    """Build dwelling prompts (and optionally the lived-experience Excel).

    Returns ``(exit_code, lived_experience_excel_path_or_None)``.

    ``allowed_dimension_slugs`` restricts generation to dimensions applicable for
    the caller's selected analytical layers (from ``svarupa_concept_layer``).
    """
    if generate_lived_experience and combined_only:
        raise ValueError(
            "--generate-lived-experience requires per-concept prompt files; remove --combined-only."
        )
    if generate_lived_experience:
        _require_openpyxl()

    resolved_output_dir = output_dir or _default_output_dir()
    resolved_log_file: Path | None = None
    if setup_logging:
        if not no_log_file:
            resolved_log_file = log_file or (resolved_output_dir / "generation.log")
        _setup_logging(verbose=verbose, log_file=resolved_log_file)

    reset_step2_bedrock_client()
    cfg = _cfg()

    logger.info(
        "Starting prompt generation | output_dir=%s dimension_filter=%r concept_filter=%r "
        "allowed_dimensions=%s limit=%s combined_only=%s bedrock_model=%r bedrock_region=%r",
        resolved_output_dir,
        dimension,
        concept,
        sorted(allowed_dimension_slugs) if allowed_dimension_slugs is not None else None,
        limit,
        combined_only,
        model_id,
        region,
    )
    logger.info(
        "Connecting to MySQL | user=%s host=%s database=%s",
        cfg["user"],
        cfg["host"],
        cfg["database"],
    )

    conn = pymysql.connect(
        host=str(cfg["host"]),
        port=int(cfg["port"]),
        user=str(cfg["user"]),
        password=str(cfg["password"]),
        database=str(cfg["database"]),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
    )
    try:
        with conn.cursor() as cur:
            records = _fetch_concepts(
                cur,
                cfg,
                dimension=dimension,
                concept=concept,
                allowed_dimension_slugs=allowed_dimension_slugs,
            )
            logger.info("Fetched %s concept(s) from database", len(records))
    finally:
        conn.close()

    if limit is not None:
        records = records[: max(0, limit)]
        logger.info("Applied concept limit=%s → %s concept(s)", limit, len(records))

    if not records:
        logger.error("No concepts found matching filters.")
        return 1, None

    combined_path, written, succeeded, failed = _write_prompts(
        records,
        resolved_output_dir,
        combined_only=combined_only,
        model_id=model_id,
        region=region,
        max_workers=max_workers,
    )

    logger.info(
        "Generation complete | concepts=%s succeeded=%s failed=%s per_concept_files=%s combined_file=%s",
        len(records),
        succeeded,
        failed,
        len(written),
        combined_path,
    )

    le_failed = 0
    excel_path: Path | None = None
    if generate_lived_experience:
        excel_path = lived_experience_output or (
            resolved_output_dir / "dimension_dwelling_lived_experience.xlsx"
        )
        checkpoint_path = resolved_output_dir / "lived_experience_checkpoint.json"
        le_model_id = lived_experience_model_id or model_id
        logger.info(
            "Starting lived-experience generation | concepts=%s output=%s checkpoint=%s "
            "bedrock_model=%r bedrock_region=%r resume=%s",
            len(written),
            excel_path,
            checkpoint_path,
            le_model_id,
            region,
            resume,
        )
        le_succeeded, le_failed = _generate_lived_experiences(
            written,
            output_path=excel_path,
            checkpoint_path=checkpoint_path,
            resume=resume,
            model_id=le_model_id,
            region=region,
            max_tokens=lived_experience_max_tokens,
            max_workers=max_workers,
        )
        logger.info(
            "Lived-experience generation complete | concepts=%s succeeded=%s failed=%s output=%s",
            len(written),
            le_succeeded,
            le_failed,
            excel_path,
        )

    return (0 if failed == 0 and le_failed == 0 else 1), excel_path


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Generate dimension-dwelling lived-experience prompts for all concepts."
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="Output directory (default: ./prompts/dd-mm-yyyy under current working directory).",
    )
    parser.add_argument(
        "--dimension",
        type=str,
        default="",
        help="Optional filter: only concepts in this dimension.",
    )
    parser.add_argument(
        "--concept",
        type=str,
        default="",
        help="Optional filter: only this concept slug (e.g. annamaya_kosha).",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Optional: process only the first N concepts after filters.",
    )
    parser.add_argument(
        "--combined-only",
        action="store_true",
        help="Write only all_prompts.txt (still writes per-concept files by default).",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        help="Enable DEBUG logging for operational details.",
    )
    parser.add_argument(
        "--log-file",
        type=Path,
        default=None,
        help=(
            "Write operational logs to this file (default: <output-dir>/generation.log). "
            "Bedrock API request prompts are logged; assembled prompt files are not. "
            "Use --no-log-file for console only."
        ),
    )
    parser.add_argument(
        "--no-log-file",
        action="store_true",
        help="Do not write a generation.log file under the output directory.",
    )
    parser.add_argument(
        "--model-id",
        type=str,
        default="",
        help=(
            "Bedrock model / inference profile for Step 2 example generation "
            "(default: BEDROCK_MODEL_ID or us.anthropic.claude-opus-4-7)."
        ),
    )
    parser.add_argument(
        "--region",
        type=str,
        default="",
        help="AWS region for Bedrock (default: AWS_REGION or us-west-2).",
    )
    parser.add_argument(
        "--generate-lived-experience",
        action="store_true",
        help=(
            "After writing per-(concept, status) prompt files, send each prompt through the LLM "
            "and store the generated lived-experience statements (20 per status: 10 Western + "
            "10 Asian) in an Excel workbook. Requires per-concept prompt files (incompatible with "
            "--combined-only)."
        ),
    )
    parser.add_argument(
        "--lived-experience-output",
        type=Path,
        default=None,
        help=(
            "Excel output path for lived-experience statements "
            "(default: <output-dir>/dimension_dwelling_lived_experience.xlsx)."
        ),
    )
    parser.add_argument(
        "--lived-experience-model-id",
        type=str,
        default="",
        help="Bedrock model / inference profile for lived-experience generation (default: --model-id).",
    )
    parser.add_argument(
        "--lived-experience-max-tokens",
        type=int,
        default=LIVED_EXPERIENCE_MAX_TOKENS,
        help=f"Max tokens for lived-experience generation (default: {LIVED_EXPERIENCE_MAX_TOKENS}).",
    )
    parser.add_argument(
        "--max-workers",
        type=int,
        default=DEFAULT_MAX_WORKERS,
        help=(
            "Concurrent Bedrock calls for Step 2 and lived-experience generation "
            f"(default: {DEFAULT_MAX_WORKERS}). Raise for throughput, lower if throttled."
        ),
    )
    parser.add_argument(
        "--no-resume",
        action="store_true",
        help=(
            "Ignore the lived-experience checkpoint file and regenerate every concept "
            "(only relevant with --generate-lived-experience)."
        ),
    )
    args = parser.parse_args()

    if args.generate_lived_experience and args.combined_only:
        parser.error("--generate-lived-experience requires per-concept prompt files; remove --combined-only.")
    if args.generate_lived_experience:
        try:
            _require_openpyxl()
        except RuntimeError as exc:
            parser.error(str(exc))

    exit_code, _excel = run_dimension_dwelling_generation(
        output_dir=args.output_dir,
        dimension=args.dimension.strip() or None,
        concept=args.concept.strip() or None,
        limit=args.limit,
        combined_only=args.combined_only,
        generate_lived_experience=args.generate_lived_experience,
        lived_experience_output=args.lived_experience_output,
        model_id=args.model_id.strip() or None,
        region=args.region.strip() or None,
        lived_experience_model_id=args.lived_experience_model_id.strip() or None,
        lived_experience_max_tokens=args.lived_experience_max_tokens,
        max_workers=args.max_workers,
        resume=not args.no_resume,
        verbose=args.verbose,
        log_file=args.log_file,
        no_log_file=args.no_log_file,
        setup_logging=True,
    )
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
